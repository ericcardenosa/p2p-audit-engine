#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""P2P Audit Engine | Version 1.1.0 | Synthetic portfolio demonstration.

Independent, reproducible purchase-to-pay analytics in Excel.
Data: 100% synthetic four-source demo (purchase orders, invoices, projects, CECOs).
No bank statements, approvals logs, receipt confirmations, or general ledger: the
engine DOES NOT conclude that a payment was executed or an expense was misstated.

Usage:
  python audit_engine.py --input-dir inputs --output output/P2P_AUDIT_2026.xlsx
  python audit_engine.py --input-dir inputs --output output/P2P_AUDIT_2026.xlsx --tolerance 1.00

Dependency: openpyxl (pip install -r requirements.txt).
"""
from __future__ import annotations

import argparse
import calendar
import collections
import hashlib
import json
import logging
import re
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

VERSION = "1.1.0"
MONTHS = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
NAVY, INK, BLUE, LIGHT, WHITE = "#101E32", "#172B47", "#2463A1", "#EFF4F9", "#FFFFFF"
AQUA, GREEN, AMBER, RED, MUTED = "#0A8D95", "#167B62", "#B6781B", "#AD3443", "#63758B"
EUR = '#,##0.00 "€";[Red](#,##0.00 "€");–'
DATEFMT = "dd/mm/yyyy"
PO_RE = re.compile(r"(?<![A-Za-z0-9])(O\d{6,})(?!\d)", re.I)
PROJ_RE = re.compile(r"\b(XPR26-\d{5})\b", re.I)


def safe_str(val: Any) -> str:
    return "" if val is None else str(val).strip()


def amount(val: Any) -> float:
    if isinstance(val, (int, float)) and not isinstance(val, bool):
        return round(float(val), 2)
    s = safe_str(val).replace("€", "").replace("\u00a0", "").replace(" ", "")
    if not s:
        return 0.0
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return round(float(s), 2)
    except ValueError:
        return 0.0


def date_from_excel(val: Any) -> datetime | None:
    if isinstance(val, datetime):
        return val.replace(hour=0, minute=0, second=0, microsecond=0)
    if isinstance(val, (int, float)):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(val))).replace(hour=0, minute=0, second=0, microsecond=0)
        except (OverflowError, ValueError):
            return None
    if isinstance(val, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(val.strip()[:10], fmt)
            except ValueError:
                pass
    return None


def dt_text(v: datetime | None) -> str:
    return v.strftime("%d/%m/%Y") if v else ""


def cell_safe(val: Any) -> Any:
    """Avoid Excel formula injection from free-text source cells."""
    if isinstance(val, str) and val.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + val
    return val


def norm(s: Any) -> str:
    return re.sub(r"\s+", "", safe_str(s)).lower()


def extract_refs(s: Any) -> list[str]:
    return list(dict.fromkeys(x.upper() for x in PO_RE.findall(safe_str(s))))


def project_id(s: Any) -> str:
    m = PROJ_RE.search(safe_str(s))
    return m.group(1).upper() if m else ""


def label_or_blank(obj: dict, key: str) -> str:
    return safe_str(obj.get(key))


def input_file(directory: Path, prefix: str) -> Path:
    found = sorted(directory.glob(prefix + "*.xlsx"))
    if len(found) != 1:
        raise ValueError(f"Se esperaba un archivo '{prefix}*.xlsx' en {directory}; encontrados: {len(found)}")
    return found[0]


def read_source(path: Path, columns: int, start_row: int, end_col: str, max_rows: int = 12000) -> list[list[Any]]:
    """Read the first worksheet from start_row, keeping the first `columns` cells of non-empty rows."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = []
        for r in ws.iter_rows(min_row=start_row, max_row=max_rows, values_only=True):
            r = list(r[:columns]) + [None] * max(0, columns - len(r))
            if any(x is not None and x != "" for x in r):
                rows.append(r)
        return rows
    finally:
        wb.close()


@dataclass
class Sources:
    orders: list[dict]
    invoices: list[dict]
    project_ids: set[str]
    cecos: dict[str, str]
    files: dict[str, str]


def load_data(directory: Path) -> Sources:
    paths = {k: input_file(directory, prefix) for k, prefix in {
        "po": "X_DEMO_Pedidos_Compra", "invoice": "X_DEMO_Facturas_Compras",
        "project": "X_DEMO_Proyectos", "ceco": "X_DEMO_CECOs",
    }.items()}
    po_rows = read_source(paths["po"], 21, 6, "U")
    invoice_rows = read_source(paths["invoice"], 21, 6, "U")
    proj_rows = read_source(paths["project"], 9, 6, "I")
    ceco_rows = read_source(paths["ceco"], 1, 2, "A", max_rows=1500)
    orders = []
    for ix, r in enumerate(po_rows, 6):
        po = safe_str(r[0]).upper()
        if not PO_RE.fullmatch(po):
            continue  # Includes synthetic export footers.
        orders.append(dict(po=po, row=ix, date=date_from_excel(r[1]), vendor=safe_str(r[4]),
                           description=safe_str(r[5]), tags=safe_str(r[6]), account=safe_str(r[7]),
                           project=safe_str(r[8]), subtotal=amount(r[14]), total=amount(r[15]),
                           erp_invoiced=amount(r[16]), approval=safe_str(r[19]), state=safe_str(r[20])))
    invoices = []
    for ix, r in enumerate(invoice_rows, 6):
        num = safe_str(r[1])
        if not num.startswith("FX26-"):
            continue
        refs = extract_refs(r[6])
        invoices.append(dict(num=num, row=ix, date=date_from_excel(r[0]), acct_date=date_from_excel(r[3]),
                             due=date_from_excel(r[4]), vendor=safe_str(r[5]), description=safe_str(r[6]),
                             tags=safe_str(r[7]), account=safe_str(r[8]), payment_method=safe_str(r[9]),
                             project=safe_str(r[10]), subtotal=amount(r[11]), iva=amount(r[12]),
                             withholding=amount(r[13]), total=amount(r[16]), paid=amount(r[17]),
                             outstanding=amount(r[18]), state=safe_str(r[19]), payment_date=date_from_excel(r[20]),
                             refs=refs))
    project_ids = {project_id(r[0]) for r in proj_rows if r and project_id(r[0])}
    cecos = {norm(r[0]):safe_str(r[0]) for r in ceco_rows if r and norm(r[0]) not in ("cecos", "cecosobligatorios")}
    if len({p["po"] for p in orders}) != len(orders):
        raise ValueError("Se han detectado números PO duplicados; corrige las fuentes antes de emitir el informe.")
    if not orders or not invoices or not project_ids or not cecos:
        raise ValueError("Archivos sin datos suficientes: verifica los cuatro ficheros de entrada.")
    return Sources(orders, invoices, project_ids, cecos, {k:v.name for k,v in paths.items()})


def ceco_state(tags: str, catalog: dict[str, str]) -> tuple[str, str]:
    tokens = re.findall(r"[A-Za-zÁ-ÿ]+_x\d{3}", safe_str(tags).lower(), flags=re.I)
    if not tokens:
        return ("NO INFORMADO", "")
    matches = [catalog[norm(t)] for t in tokens if norm(t) in catalog]
    if matches:
        return ("VALIDO", ", ".join(dict.fromkeys(matches)))
    return ("NO CATALOGADO", ", ".join(dict.fromkeys(tokens)))


def add_issue(issues: list[dict], rule: str, level: str, kind: str, ref: str,
              vendor: str, evidence: str, impact: float = 0.0, source_row: int = 0):
    fingerprint = f"{rule}|{kind}|{ref}|{source_row}"
    issue_id = "EX-" + hashlib.sha1(fingerprint.encode("utf-8")).hexdigest()[:10].upper()
    issues.append(dict(id=issue_id, level=level, rule=rule, kind=kind, ref=ref, vendor=vendor,
                       evidence=evidence, impact=round(abs(impact), 2), status="Pendiente",
                       owner="", note="", source_row=source_row))


def analyze(src: Sources, tolerance: float) -> dict[str, Any]:
    po_by_id = {o["po"]:o for o in src.orders}
    link = collections.defaultdict(list)
    no_po = []
    issues = []
    invoices_summary = []
    cutoff = []
    ceco_issues = []
    possible_duplicates = []
    duplicate_bucket = collections.defaultdict(list)
    exact_ids = collections.defaultdict(list)

    for f in src.invoices:
        valid_refs = [k for k in f["refs"] if k in po_by_id]
        invalid_refs = [k for k in f["refs"] if k not in po_by_id]
        for ref in valid_refs:
            link[ref].append(f)
        invoice_flag = "VINCULADA" if valid_refs else "SIN REFERENCIA PO"
        if invalid_refs:
            invoice_flag = "PO INEXISTENTE"
            add_issue(issues, "INV_PO_NOT_FOUND", "Alta", "FACTURA", f["num"], f["vendor"],
                      "La descripción cita PO no presente en el catálogo: " + ", ".join(invalid_refs), f["total"], f["row"])
        if not valid_refs:
            no_po.append(f)
        st, parsed_ceco = ceco_state(f["tags"], src.cecos)
        pid = project_id(f["project"])
        project_state = "VALIDO" if pid in src.project_ids else ("NO INFORMADO" if not f["project"] else "NO CATALOGADO")
        if st != "VALIDO":
            ceco_issues.append(("FACTURA", f["num"], f["vendor"], f["tags"], st, f["project"], project_state, f["total"]))
            if st == "NO CATALOGADO":
                add_issue(issues, "INV_CECO_INVALID", "Media", "FACTURA", f["num"], f["vendor"],
                          "Etiqueta CECO no encontrada en el catálogo: " + f["tags"], f["total"], f["row"])
        if project_state == "NO CATALOGADO":
            add_issue(issues, "INV_PROJECT_INVALID", "Media", "FACTURA", f["num"], f["vendor"],
                      "Proyecto de factura no localizado en catálogo: " + f["project"], f["total"], f["row"])
        if abs(f["subtotal"] + f["iva"] - f["withholding"] - f["total"]) > tolerance:
            add_issue(issues, "INV_ARITHMETIC", "Alta", "FACTURA", f["num"], f["vendor"],
                      "Subtotal + IVA - retenciones ≠ total. Verificar recargo/equivalencia/excepciones.",
                      f["subtotal"] + f["iva"] - f["withholding"] - f["total"], f["row"])
        if abs(f["total"] - f["paid"] - f["outstanding"]) > tolerance:
            add_issue(issues, "INV_OPEN_BALANCE", "Alta", "FACTURA", f["num"], f["vendor"],
                      "Total - pagado - pendiente difiere; revisar campos de exportación.",
                      f["total"] - f["paid"] - f["outstanding"], f["row"])
        if f["date"] and f["acct_date"] and (f["date"].year, f["date"].month) != (f["acct_date"].year,f["acct_date"].month):
            cutoff.append(("FACTURA", f["num"], dt_text(f["date"]), dt_text(f["acct_date"]),
                           "Emisión / contabilización en meses distintos", f["total"], f["vendor"]))
            add_issue(issues, "INV_BOOKING_MONTH", "Media", "FACTURA", f["num"], f["vendor"],
                      "Mes emisión distinto de mes contable; requiere revisión de devengo.", f["total"], f["row"])
        duplicate_bucket[(norm(f["vendor"]), round(f["total"],2), dt_text(f["date"]))].append(f)
        exact_ids[(norm(f["vendor"]), norm(f["num"]))].append(f)
        invoices_summary.append([f["num"], dt_text(f["date"]), dt_text(f["acct_date"]), f["vendor"], f["description"],
                                 ", ".join(valid_refs) or ", ".join(invalid_refs), invoice_flag,
                                 f["subtotal"], f["iva"], f["total"], f["paid"], f["outstanding"],
                                 f["state"], f["project"], project_state, st, f["account"], f["row"]])

    for group in exact_ids.values():
        if len(group) > 1:
            for f in group:
                possible_duplicates.append(("ID REPETIDO",f))
                add_issue(issues, "DUP_VENDOR_INVOICE", "Alta", "FACTURA", f["num"], f["vendor"],
                          "Se repite combinación proveedor + número de factura.", f["total"], f["row"])
    for group in duplicate_bucket.values():
        if len(group) > 1:
            for f in group:
                possible_duplicates.append(("COINCIDENCIA INDICIARIA", f))
                add_issue(issues, "DUP_SAME_AMOUNT_DATE", "Media", "FACTURA", f["num"], f["vendor"],
                          "Mismo proveedor, fecha e importe que otra factura; no prueba duplicidad.", f["total"], f["row"])

    reconciliation = []
    po_details = []
    no_invoice = []
    link_rows = []
    for o in src.orders:
        invs = link.get(o["po"], [])
        allocated = sum(f["total"] / max(1,len([k for k in f["refs"] if k in po_by_id])) for f in invs)
        diff = round(o["total"] - allocated, 2)
        if invs:
            status = "SOBRE PEDIDO" if diff < -tolerance else "PARCIAL / DIFERENCIA" if diff > tolerance else "CONCILIADO"
        elif o["erp_invoiced"] and abs(o["erp_invoiced"]) > tolerance:
            status = "ERP FACTURADO / SIN ENLACE"
        elif o["approval"].casefold() == "cancelado":
            status = "CANCELADO / SIN FACTURA"
        else:
            status = "SIN FACTURA LOCALIZADA"
        st, ce = ceco_state(o["tags"], src.cecos)
        pid = project_id(o["project"])
        project_state = "VALIDO" if pid in src.project_ids else ("NO INFORMADO" if not o["project"] else "NO CATALOGADO")
        reconciliation.append([o["po"], dt_text(o["date"]), o["vendor"], o["description"], o["total"], len(invs),
                               round(allocated,2), diff, o["erp_invoiced"], status, o["approval"], o["state"],
                               ", ".join(f["num"] for f in invs), o["project"], project_state, st, ce, o["row"]])
        po_details.append([o["po"],dt_text(o["date"]),o["vendor"],o["description"],o["subtotal"],o["total"],
                           o["erp_invoiced"],o["approval"],o["state"],o["project"],project_state,
                           o["tags"],st,o["account"],len(invs),status,o["row"]])
        if not invs:
            no_invoice.append([o["po"],dt_text(o["date"]),o["vendor"],o["total"],o["erp_invoiced"],
                               o["approval"],o["state"],status,o["project"],o["tags"],o["row"]])
            if status == "ERP FACTURADO / SIN ENLACE":
                add_issue(issues,"PO_NO_TRACEABLE_INVOICE","Media","PEDIDO",o["po"],o["vendor"],
                          "ERP muestra facturado pero no se localiza factura con referencia PO.",o["erp_invoiced"],o["row"])
            elif status == "SIN FACTURA LOCALIZADA":
                add_issue(issues,"PO_INVOICE_NOT_FOUND","Media","PEDIDO",o["po"],o["vendor"],
                          "No consta factura referenciada; valorar estado, servicio recibido y accruals.",o["total"],o["row"])
        if invs and diff < -tolerance:
            add_issue(issues,"PO_INVOICED_OVER","Alta","PEDIDO",o["po"],o["vendor"],
                      f"Facturación identificada excede PO en {abs(diff):,.2f} EUR; comprobar ampliaciones/abonos.",abs(diff),o["row"])
        if invs and o["approval"].casefold()=="cancelado":
            add_issue(issues,"PO_CANCELLED_BILLED","Alta","PEDIDO",o["po"],o["vendor"],
                      "Pedido cancelado con facturas referenciadas; verificar causa.",allocated,o["row"])
        if st == "NO CATALOGADO":
            add_issue(issues,"PO_CECO_INVALID","Media","PEDIDO",o["po"],o["vendor"],
                      "Etiqueta CECO no catalogada: " + o["tags"],o["total"],o["row"])
        if project_state == "NO CATALOGADO":
            add_issue(issues,"PO_PROJECT_INVALID","Media","PEDIDO",o["po"],o["vendor"],
                      "Proyecto PO no catalogado: " + o["project"],o["total"],o["row"])
        if st != "VALIDO":
            ceco_issues.append(("PEDIDO", o["po"], o["vendor"], o["tags"], st, o["project"], project_state, o["total"]))
        for f in invs:
            link_rows.append((o["po"],f["num"],o["vendor"],f["vendor"],o["date"],f["date"],f["total"]))
            if o["date"] and f["date"]:
                if f["date"] < o["date"]:
                    cutoff.append(("PO-FACTURA", o["po"]+" / "+f["num"],dt_text(o["date"]), dt_text(f["date"]),
                                   "Factura emitida antes de PO (verificar autorización)", f["total"], f["vendor"]))
                    add_issue(issues,"INV_BEFORE_PO","Media","PEDIDO",o["po"],o["vendor"],
                              "Factura " + f["num"] + " emitida antes de la creación de PO.",f["total"],o["row"])
                elif (o["date"].year,o["date"].month)!=(f["date"].year,f["date"].month):
                    cutoff.append(("PO-FACTURA", o["po"]+" / "+f["num"],dt_text(o["date"]),dt_text(f["date"]),
                                   "PO y factura en distinto mes (no demuestra error de corte)",f["total"],f["vendor"]))
            if norm(o["vendor"]) != norm(f["vendor"]):
                add_issue(issues,"PO_VENDOR_DIFFERENT","Media","PEDIDO",o["po"],o["vendor"],
                          f"Proveedor de la PO difiere del de factura {f['num']}.",f["total"],o["row"])
    # Invariant: invoice amount across PO links must be allocated only once.
    allocated_total = round(sum(r[6] for r in reconciliation),2)
    matched_invoice_total = round(sum(f["total"] for f in src.invoices if any(k in po_by_id for k in f["refs"])),2)
    if abs(allocated_total-matched_invoice_total)>max(tolerance,0.02):
        raise AssertionError(f"Error en conciliación 1:N: {allocated_total=} vs {matched_invoice_total=}")

    # Unique exceptions - stable IDs; deterministic sorting independent of input order.
    unique = {i["id"]:i for i in issues}
    issues = sorted(unique.values(),key=lambda x:({"Alta":0,"Media":1,"Baja":2}.get(x["level"],3),x["rule"],x["ref"],x["id"]))
    monthly = []
    for m in range(1,13):
        po_m = [o for o in src.orders if o["date"] and o["date"].month==m and o["date"].year==2026]
        inv_m = [f for f in src.invoices if f["date"] and f["date"].month==m and f["date"].year==2026]
        mm_issues = [i for i in issues if (i["kind"]=="PEDIDO" and po_by_id.get(i["ref"],{}).get("date")
                     and po_by_id[i["ref"]]["date"].month==m) or (i["kind"]=="FACTURA" and any(f["num"]==i["ref"] and f["date"] and f["date"].month==m for f in inv_m))]
        monthly.append([MONTHS[m-1],len(po_m),round(sum(o["total"] for o in po_m),2),len(inv_m),round(sum(f["total"] for f in inv_m),2),
                        len(mm_issues),sum(1 for i in mm_issues if i["level"]=="Alta")])
    result = dict(src=src, reconciliation=reconciliation, po_details=po_details, invoice_details=invoices_summary,
                  no_invoice=no_invoice, no_po=no_po, duplicates=possible_duplicates,
                  cutoff=cutoff, ceco_issues=ceco_issues, issues=issues, monthly=monthly,
                  allocated_total=allocated_total, matched_invoice_total=matched_invoice_total,
                  link_rows=link_rows, tolerance=tolerance)
    return result


# --------------------------------------------------------------------------- #
# Excel rendering (openpyxl)                                                  #
# --------------------------------------------------------------------------- #
FONT_DISPLAY = "Aptos Display"
FONT_BODY = "Aptos"


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_color.lstrip("#"), end_color=hex_color.lstrip("#"))


def _font(size: float = 10, bold: bool = False, color: str = INK, name: str = FONT_BODY, italic: bool = False) -> Font:
    return Font(name=name, size=size, bold=bold, italic=italic, color=color.lstrip("#"))


def col_letter(n: int) -> str:
    return get_column_letter(n)


def _style_range(ws, ref: str, fill: str | None = None, font: Font | None = None, align: Alignment | None = None):
    for row in ws[ref]:
        for c in row:
            if fill:
                c.fill = _fill(fill)
            if font:
                c.font = font
            if align:
                c.alignment = align


def style_title(ws, title: str, subtitle: str, end: str = "R"):
    ws.merge_cells(f"A1:{end}2")
    _style_range(ws, f"A1:{end}2", NAVY, _font(19, True, WHITE, FONT_DISPLAY), Alignment(vertical="center", indent=1))
    ws["A1"] = title
    ws.merge_cells(f"A3:{end}3")
    _style_range(ws, f"A3:{end}3", INK, _font(10, color="#BBD8ED"), Alignment(vertical="center", indent=1))
    ws["A3"] = subtitle
    for r, h in ((1, 25), (2, 25), (3, 23)):
        ws.row_dimensions[r].height = h


_TABLE_NAMES: set[str] = set()


def grid_sheet(wb, name: str, title: str, subtitle: str, headers: list[str], rows: list[list[Any]],
               widths: list[int] | None = None, money_cols: tuple[int, ...] = (), caution_col: int | None = None):
    ws = wb.create_sheet(name)
    nc = len(headers)
    end = col_letter(nc)
    style_title(ws, title, subtitle, end)
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=5, column=j, value=h)
        c.fill = _fill(BLUE)
        c.font = _font(10, True, WHITE)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[5].height = 33
    last = 5
    if rows:
        sanitized = [[cell_safe(x) for x in row] for row in rows]
        assert all(len(row) == nc for row in sanitized), (name, nc, [len(x) for x in sanitized[:3]])
        for i, row in enumerate(sanitized, 6):
            for j, v in enumerate(row, 1):
                c = ws.cell(row=i, column=j, value=v)
                c.font = _font(10)
        last = 5 + len(rows)
        if caution_col is not None:
            col = col_letter(caution_col)
            ws.conditional_formatting.add(f"{col}6:{col}{last}",
                                          FormulaRule(formula=[f'ISNUMBER(SEARCH("SIN ",{col}6))'], fill=_fill("#FFF1E6")))
    for idx in range(1, nc + 1):
        ws.column_dimensions[col_letter(idx)].width = widths[idx - 1] if widths else 18
    for n in money_cols:
        for r in range(6, last + 1):
            ws.cell(row=r, column=n).number_format = EUR
    ws.freeze_panes = "A6"
    if last >= 6:
        tname = "T" + re.sub(r"[^A-Za-z0-9]", "", name)[:22] + str(nc)
        while tname in _TABLE_NAMES:
            tname += "x"
        _TABLE_NAMES.add(tname)
        tab = Table(displayName=tname, ref=f"A5:{end}{last}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(tab)
    return ws


def metric(ws, label_range: str, value_range: str, label: str, value: Any, numeric_format: str | None = None):
    ws.merge_cells(label_range)
    ws.merge_cells(value_range)
    _style_range(ws, label_range, LIGHT, _font(10, True, MUTED), Alignment(vertical="center", indent=1))
    _style_range(ws, value_range, WHITE, _font(20, True, NAVY, FONT_DISPLAY), Alignment(horizontal="right", vertical="center", indent=1))
    ws[label_range.split(":")[0]] = label
    top = value_range.split(":")[0]
    ws[top] = value
    if numeric_format:
        ws[top].number_format = numeric_format


def build_dashboard(wb, result: dict):
    ws = wb.create_sheet("01_Dashboard", 0)
    for col in "ABCDEFGHIJKLMNO":
        ws.column_dimensions[col].width = 13
    ws.column_dimensions["A"].width = 17
    ws.column_dimensions["D"].width = 17
    ws.column_dimensions["F"].width = 17
    ws.column_dimensions["H"].width = 17
    style_title(ws, "P2P  /  AUDIT COMMAND CENTER", "SOCIEDAD X, S.L.  •  DEMO SINTÉTICA  •  FY 2026  •  v" + VERSION, "O")
    ws.merge_cells("B5:C5")
    ws["B5"] = "PERIODO DE ANÁLISIS"
    _style_range(ws, "B5:C5", NAVY, _font(10, True, WHITE), Alignment(vertical="center", indent=1))
    ws.merge_cells("D5:E5")
    ws["D5"] = "Todos"
    _style_range(ws, "D5:E5", "#DEF3F1", _font(11, True, GREEN), Alignment(vertical="center", indent=1))
    dv = DataValidation(type="list", formula1='"' + ",".join(["Todos"] + MONTHS) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("D5")
    ws.row_dimensions[5].height = 26
    ws.merge_cells("G5:O5")
    ws["G5"] = "El selector actualiza importes y volúmenes de cada mes. Las alertas son indicios, no conclusiones."
    ws["G5"].font = _font(10, color=MUTED, italic=True)
    m = "'11_Resumen_Mensual'"

    def sf(col):
        return f'=IF($D$5="Todos",SUM({m}!{col}6:{col}17),SUMIF({m}!$A$6:$A$17,$D$5,{m}!{col}$6:{col}$17))'

    metric(ws, "B7:E7", "B8:E9", "PEDIDOS REGISTRADOS", sf("B"), "0")
    metric(ws, "G7:J7", "G8:J9", "IMPORTE TOTAL POs", sf("C"), EUR)
    metric(ws, "L7:O7", "L8:O9", "FACTURAS REGISTRADAS", sf("D"), "0")
    metric(ws, "B11:E11", "B12:E13", "IMPORTE TOTAL FACTURAS", sf("E"), EUR)
    metric(ws, "G11:J11", "G12:J13", "EXCEPCIONES DETECTADAS", sf("F"), "0")
    metric(ws, "L11:O11", "L12:O13", "EXCEPCIONES PRIORIDAD ALTA", sf("G"), "0")
    for r in (7, 11):
        ws.row_dimensions[r].height = 26
    ws.merge_cells("B15:O15")
    ws["B15"] = "ANÁLISIS MENSUAL  |  VOLUMEN OPERATIVO Y VALOR (€)"
    _style_range(ws, "B15:O15", INK, _font(11, True, WHITE), Alignment(vertical="center", indent=1))
    ws.row_dimensions[15].height = 27
    hd = ["Mes", "Nº POs", "Pedidos (€)", "Nº facturas", "Facturas (€)", "Alertas", "Altas"]
    for j, h in enumerate(hd, 2):
        c = ws.cell(row=17, column=j, value=h)
        c.fill = _fill(BLUE)
        c.font = _font(10, True, WHITE)
    ws.row_dimensions[17].height = 27
    for i in range(12):
        r = i + 18
        mr = i + 6
        for c, sc in enumerate("ABCDEFG", start=2):
            cell = ws.cell(row=r, column=c, value=f"={m}!{sc}{mr}")
            cell.font = _font(10)
    for r in range(18, 30):
        ws.cell(row=r, column=4).number_format = EUR
        ws.cell(row=r, column=6).number_format = EUR
    ms = wb["11_Resumen_Mensual"]
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = "Tendencia mensual: pedidos vs. facturas"
    chart.y_axis.numFmt = '#,##0'
    chart.add_data(Reference(ms, min_col=3, min_row=5, max_row=17), titles_from_data=True)
    chart.add_data(Reference(ms, min_col=5, min_row=5, max_row=17), titles_from_data=True)
    chart.set_categories(Reference(ms, min_col=1, min_row=6, max_row=17))
    for series, color, label in zip(chart.series, ("1F5C8A", "E8793A"), ("Pedidos (€)", "Facturas (€)")):
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.solidFill = color
        series.tx = SeriesLabel(v=label)
    chart.height = 9
    chart.width = 17
    ws.add_chart(chart, "I17")
    ws.merge_cells("B32:H32")
    ws["B32"] = "LECTURA CORRECTA DEL DASHBOARD"
    _style_range(ws, "B32:H32", INK, _font(10, True, WHITE), Alignment(vertical="center", indent=1))
    explanations = [
        "Los totales de POs y facturas son series independientes y no representan cobertura monetaria.",
        "Las excepciones de fecha se revisan, pero no prueban errores de corte contable.",
        "El campo 'Pagado' procede de un export, no de confirmación bancaria ni de libro mayor.",
        "Los números con signo negativo (abonos) se conservan en las sumas.",
        "El historial de decisiones se registra en 05_Excepciones antes de volver a generar el fichero.",
    ]
    for idx, s in enumerate(explanations, start=33):
        ws.merge_cells(f"B{idx}:O{idx}")
        ws[f"B{idx}"] = "•  " + s
        ws[f"B{idx}"].font = _font(10, color=MUTED)
        ws.row_dimensions[idx].height = 20
    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False
    return ws


def build_workbook(result:dict, output:Path) -> dict:
    wb=Workbook()
    wb.remove(wb.active)
    _TABLE_NAMES.clear()
    rec=result["reconciliation"]
    grid_sheet(wb,"02_Reconciliacion","CONCILIACIÓN PEDIDOS ↔ FACTURAS",
               "Cruce 1:N por referencia PO en descripción; importes al total con IVA, no confundir con gasto devengado",
               ["PO","Fecha PO","Proveedor","Descripción","Total PO €","Nº facturas","Facturado enlazado €","Dif. PO - fact. €","Facturado ERP €","Resultado","Aprobación PO","Estado PO","Facturas vinculadas","Proyecto","Estado proyecto","Estado CECO","CECO identificado","Fila origen"],
               rec,[18,15,31,45,20,15,23,22,22,29,19,18,52,40,20,18,31,13], (5,7,8,9),10)
    grid_sheet(wb,"03_Facturas","UNIVERSO DE FACTURAS",
               "Todas las facturas de Compras; mantener abonos y facturas sin referencia PO; fecha contable vs emisión",
               ["Factura","Emisión","Fecha contable","Proveedor","Descripción","PO citada","Vínculo","Subtotal €","IVA €","Total €","Pagado €","Pendiente €","Estado","Proyecto","Val. proyecto","Val. CECO","Cuenta","Fila origen"],
               result["invoice_details"],[21,15,15,32,53,20,23,18,18,18,18,18,19,40,20,20,32,13],(8,9,10,11,12),7)
    grid_sheet(wb,"04_Pedidos","UNIVERSO DE PEDIDOS",
               "Fuente completa de pedidos. La presencia de factura no sustituye evidencia de recepción ni autorización.",
               ["PO","Fecha","Proveedor","Descripción","Subtotal €","Total €","Facturado ERP €","Aprobación","Estado","Proyecto","Val. proyecto","Etiquetas","Val. CECO","Cuenta","Nº facturas","Vínculo","Fila origen"],
               result["po_details"],[18,15,34,48,18,18,21,20,20,39,20,39,19,35,15,29,13],(5,6,7),16)
    issues=result["issues"]
    issue_rows=[[i[x] for x in ("id","level","rule","kind","ref","vendor","evidence","impact","status","owner","note","source_row")] for i in issues]
    es=grid_sheet(wb,"05_Excepciones","REGISTRO DE EXCEPCIONES / WORKPAPERS",
             "Cada alerta es una hipótesis de riesgo; el auditor debe documentar evidencia, criterio y resolución.",
             ["ID estable","Prioridad","Control","Origen","Referencia","Proveedor","Motivo / evidencia pendiente","Impacto indicativo €","Estado revisión","Responsable","Conclusión / notas","Fila origen"],
             issue_rows,[22,15,26,17,23,33,82,24,19,22,50,13],(8,),2)
    if issues:
        dv=DataValidation(type="list",formula1='"Pendiente,En investigación,Justificada,Corregida,No aplica"',allow_blank=True)
        es.add_data_validation(dv); dv.add(f"I6:I{5+len(issues)}")
        es.conditional_formatting.add(f"B6:B{5+len(issues)}",FormulaRule(formula=['B6="Alta"'],fill=_fill("#FADDE0"),font=Font(color=RED.lstrip("#"),bold=True)))
    grid_sheet(wb,"06_Sin_Factura","PEDIDOS SIN FACTURA VINCULADA",
             "Diferenciar 'facturado ERP sin referencia' de 'sin factura localizada' y de pedidos cancelados.",
             ["PO","Fecha","Proveedor","Total PO €","Facturado ERP €","Aprobación","Estado","Clasificación","Proyecto","Etiquetas","Fila origen"],
             result["no_invoice"],[20,16,34,20,23,18,18,33,44,48,13],(4,5),8)
    no_po_rows=[]
    for f in result["no_po"]:
        no_po_rows.append([f["num"],dt_text(f["date"]),f["vendor"],f["description"],f["total"],
                           ", ".join(f["refs"]),f["state"],f["project"],f["row"]])
    grid_sheet(wb,"07_Sin_PO","FACTURAS SIN PO ENLAZADA",
             "No todas las compras requieren PO; aplicar política de autorizaciones antes de concluir incumplimientos.",
             ["Factura","Emisión","Proveedor","Descripción","Total €","Ref PO no localizada","Estado","Proyecto","Fila origen"],
             no_po_rows,[21,15,32,62,20,26,21,45,12],(5,))
    duplicates_rows=[]
    for typ,f in result["duplicates"]:
        duplicates_rows.append([typ,f["num"],f["vendor"],dt_text(f["date"]),f["total"],f["state"],f["row"]])
    grid_sheet(wb,"08_Duplicados","DETECCIÓN DE DUPLICIDADES POTENCIALES",
             "Indicadores: proveedor+numero o proveedor+fecha+importe; requieren evidencia de facturas originales.",
             ["Tipo","Factura","Proveedor","Fecha","Total €","Estado","Fila origen"],
             duplicates_rows,[33,22,40,18,20,22,13],(5,))
    grid_sheet(wb,"09_Fechas","CONTROLES DE FECHAS / CUT-OFF",
             "Comparaciones preliminares. PO ≠ fecha de prestación; la diferencia de meses NO prueba ajuste.",
             ["Tipo","Documento(s)","Fecha inicial","Fecha comparada","Observación","Total factura €","Proveedor"],
             list(result["cutoff"]),[22,36,18,20,67,20,40],(6,))
    grid_sheet(wb,"10_CECO_Proyecto","VALIDACIONES CECO Y PROYECTO",
             "Las ausencias pueden ser admisibles; solo 'no catalogado' indica fallo de integridad referencial.",
             ["Tipo","Referencia","Proveedor","Etiquetas","Validación CECO","Proyecto","Validación proyecto","Total €"],
             list(result["ceco_issues"]),[18,23,37,53,22,46,25,22],(8,))
    monthly=result["monthly"]
    ms=grid_sheet(wb,"11_Resumen_Mensual","CUADRO DE CONTROL / EVOLUCIÓN 2026",
             "Los importes mensuales se asignan a su respectiva fecha de documento; no equivalen a caja ni a devengo.",
             ["Mes","Nº POs","Total POs €","Nº facturas","Total facturas €","Excepciones","Prioridad alta"],
             monthly,[19,17,22,19,25,18,19],(3,5))
    # Monthly data stays in A6:G17; used by Dashboard formulas.
    methods=[
        ["01", "OBJETIVO", "Evaluar riesgos del ciclo de compras con análisis automatizado de POs, facturas, CECOs y proyectos."],
        ["02", "FUENTES", ", ".join(result["src"].files.values())],
        ["03", "UNIVERSO", f"{len(result['src'].orders)} POs, {len(result['src'].invoices)} facturas, {len(result['src'].project_ids)} proyectos, {len(result['src'].cecos)} CECOs."],
        ["04", "RELACIÓN", "Referencias O###### en descripción de factura. Una PO admite varias facturas y cada factura multirreferencia se prorratea."],
        ["05", "NO MATCH", "Factura sin PO explícita ≠ factura improcedente. PO sin factura explícita ≠ gasto omitido."],
        ["06", "IMPORTE", "Diferencias de POs calculadas al total bruto con IVA; abonos negativos conservados."],
        ["07", "CORTE", "Fechas de PO, emisión y contable son pruebas preliminares. Se requiere evidencia de entrega/devengo."],
        ["08", "PAGOS", "Pagado/Pendiente procede del export de compras; no hay confirmación externa de banco."],
        ["09", "CECO", "Solo cotejo por tokens completos etiqueta _X### contra catálogo; no por subcadenas ambiguas."],
        ["10", "PROYECTO", "Cotejo por código XPR26-##### con maestro proyectos; ausencia no implica necesariamente error."],
        ["11", "DUPLICIDAD", "Búsqueda exacta o indicio (proveedor + fecha + importe). No confirma pagos duplicados."],
        ["12", "LIMITACIONES", "Sin recepción de mercancías, contratos, aprobación de PO con sello temporal, libro mayor ni bancos."],
        ["13", "EVIDENCIA", "Registros incluyen referencia y fila de origen; guardar documentos y conclusión manual por excepción."],
        ["14", "UMBRALES", f"Tolerancia de diferencia absoluta: {result['tolerance']:,.2f} €; no se ha aplicado materialidad de auditoría."],
        ["15", "PRIVACIDAD", "SOLO DATOS SINTÉTICOS. Uso para demostración/portfolio; no conectar a fuentes privadas."],
        ["16", "REGENERACIÓN", "Al ejecutar el script se reemplaza el XLSX: exportar/guardar el estado del registro de excepciones antes."],
    ]
    grid_sheet(wb,"12_Metodologia","METODOLOGÍA, ALCANCE Y LIMITACIONES",
             "Guía de lectura: alcance del procedimiento, control de calidad y limitaciones probatorias.",
             ["Paso","Área","Explicación / criterio"],methods,[12,28,110])
    catalog_rows=[["CECO",v,""] for v in result["src"].cecos.values()]+[["PROYECTO",k,""] for k in sorted(result["src"].project_ids)]
    grid_sheet(wb,"13_Catalogos","CATÁLOGOS MAESTROS",
             "CECOs y proyectos usados exclusivamente para validación referencial.",
             ["Tipo","Código / valor","Observación"],catalog_rows,[22,55,26])
    build_dashboard(wb,result)
    wb.active=0
    output.parent.mkdir(parents=True,exist_ok=True)
    wb.save(str(output))
    metrics={
       "version": VERSION,"society":"SOCIEDAD FICTICIA X, S.L.","orders":len(result["src"].orders),
       "invoices":len(result["src"].invoices),"projects":len(result["src"].project_ids),"cecos":len(result["src"].cecos),
       "matched_pos":sum(r[5]>0 for r in rec),"matched_invoices":len(result["src"].invoices)-len(result["no_po"]),
       "no_invoice_pos":len(result["no_invoice"]),"no_po_invoices":len(result["no_po"]),
       "exception_count":len(issues),"high_priority":sum(i["level"]=="Alta" for i in issues),
       "duplicity_flags":len(result["duplicates"]),"date_checks":len(result["cutoff"]),
       "invoice_total":round(sum(f["total"] for f in result["src"].invoices),2),
       "order_total":round(sum(p["total"] for p in result["src"].orders),2),
       "linked_invoice_total":result["matched_invoice_total"],"allocated_invoice_total":result["allocated_total"],
       "output":output.name,"engine":"openpyxl"
    }
    (output.parent/"audit_run_summary.json").write_text(json.dumps(metrics,ensure_ascii=False,indent=2),encoding="utf-8")
    return metrics


def main(argv=None):
    p=argparse.ArgumentParser(description="P2P Audit Engine - informe de auditoría automatizada para Sociedad X (datos sintéticos)")
    p.add_argument("--input-dir",default="inputs",type=Path,help="Carpeta que contiene los 4 archivos X_DEMO*.xlsx")
    p.add_argument("--output",default=Path("output/P2P_AUDIT_2026.xlsx"),type=Path,help="Ruta de salida del Excel")
    p.add_argument("--tolerance",default=1.0,type=float,help="Tolerancia absoluta en EUR, por defecto 1.0")
    p.add_argument("--verbose",action="store_true")
    args=p.parse_args(argv)
    logging.basicConfig(level=logging.INFO if args.verbose else logging.WARNING,format="%(levelname)s: %(message)s")
    if args.tolerance < 0: p.error("--tolerance debe ser >= 0")
    print("[1/3] Leyendo las cuatro fuentes sintéticas...")
    src=load_data(args.input_dir)
    print("[2/3] Ejecutando reglas P2P y comprobaciones cruzadas...")
    analysis=analyze(src,args.tolerance)
    print("[3/3] Generando reporte Excel profesional...")
    meta=build_workbook(analysis,args.output)
    print(json.dumps(meta,ensure_ascii=False,indent=2))


if __name__=="__main__":
    main()
