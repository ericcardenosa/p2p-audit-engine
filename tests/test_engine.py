"""Regression tests for the P2P Audit Engine (synthetic demo data).

Run from the project root:
    python -m unittest discover -s tests -v
"""
import random
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import audit_engine as ae  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

INPUTS = ROOT / "inputs"


class P2PAuditEngineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.src = ae.load_data(INPUTS)
        cls.result = ae.analyze(cls.src, tolerance=1.0)

    def test_source_volumes(self):
        self.assertEqual(len(self.src.orders), 427)
        self.assertEqual(len(self.src.invoices), 2239)
        self.assertEqual(len(self.src.project_ids), 3916)
        self.assertEqual(len(self.src.cecos), 277)

    def test_no_double_counting_in_1_to_n_allocation(self):
        self.assertAlmostEqual(self.result["allocated_total"], self.result["matched_invoice_total"], places=2)

    def test_exception_counts(self):
        issues = self.result["issues"]
        self.assertEqual(len(issues), 248)
        self.assertEqual(sum(i["level"] == "Alta" for i in issues), 81)

    def test_exception_ids_are_unique(self):
        ids = [i["id"] for i in self.result["issues"]]
        self.assertEqual(len(ids), len(set(ids)))

    def test_results_do_not_depend_on_input_order(self):
        shuffled = ae.Sources(list(self.src.orders), list(self.src.invoices), set(self.src.project_ids),
                              dict(self.src.cecos), dict(self.src.files))
        random.Random(7).shuffle(shuffled.orders)
        random.Random(11).shuffle(shuffled.invoices)
        again = ae.analyze(shuffled, tolerance=1.0)
        self.assertEqual([i["id"] for i in again["issues"]], [i["id"] for i in self.result["issues"]])

    def test_monthly_summary_ties_to_population(self):
        monthly_invoices = sum(row[3] for row in self.result["monthly"])
        dated_2026 = sum(1 for f in self.src.invoices if f["date"] and f["date"].year == 2026)
        self.assertEqual(monthly_invoices, dated_2026)
        monthly_amount = round(sum(row[4] for row in self.result["monthly"]), 2)
        expected = round(sum(f["total"] for f in self.src.invoices if f["date"] and f["date"].year == 2026), 2)
        self.assertAlmostEqual(monthly_amount, expected, places=2)

    def test_workbook_structure(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out" / "P2P_AUDIT_2026.xlsx"
            metrics = ae.build_workbook(self.result, out)
            wb = load_workbook(out)
            self.assertEqual(len(wb.sheetnames), 13)
            self.assertEqual(wb.sheetnames[0], "01_Dashboard")
            self.assertEqual(metrics["exception_count"], 248)
            self.assertTrue((out.parent / "audit_run_summary.json").exists())


if __name__ == "__main__":
    unittest.main()
